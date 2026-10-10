"""Build the Capstone Review 2 deck (desktop screenshots from the real app; numbers from reports/*.json).

    python -m scripts.build_ppt            -> demo/Capstone_Review2_DidThatTradeReallyHappen.pptx

Written so that a presenter WITHOUT technical background can explain it: one idea per slide, everyday analogies, and a
speaking script in the notes of every slide.
"""
from __future__ import annotations

import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
SHOTS = ROOT / "assets/screenshots"
OUT = ROOT / "demo/Capstone_Review2_DidThatTradeReallyHappen.pptx"
NAVY, INK, INK2, CREAM, OK, BAD, WARN, LINE, WHITE = (RGBColor(0x1F, 0x4E, 0x79), RGBColor(0x14, 0x20, 0x2E), RGBColor(0x44, 0x50, 0x5E),
                                                      RGBColor(0xF6, 0xF4, 0xEF), RGBColor(0x1D, 0x6B, 0x4F), RGBColor(0xA3, 0x28, 0x2F),
                                                      RGBColor(0x8A, 0x5A, 0x00), RGBColor(0xD9, 0xD4, 0xCA), RGBColor(0xFF, 0xFF, 0xFF))
FONT = "Segoe UI"
ACCENT, MUTED = RGBColor(0x0B, 0x5C, 0x8E), RGBColor(0x8A, 0x93, 0x9C)


def load_numbers():
    ex = json.loads((ROOT / "reports/experiments.json").read_text(encoding="utf-8"))
    e5p = ROOT / "reports/tau_validation.json"
    e5 = json.loads(e5p.read_text(encoding="utf-8")) if e5p.exists() else None
    return ex, e5


class Deck:
    def __init__(self):
        self.p = Presentation()
        self.p.slide_width, self.p.slide_height = Inches(13.333), Inches(7.5)
        self.n = 0

    SECTIONS = ["Motivation", "Problem", "Method", "Evidence", "System", "Outlook"]

    def slide(self, title, kicker=None, notes="", section=None):
        """research-talk frame: white page, section tracker, the slide's CLAIM as the title, thin rule, quiet footer."""
        s = self.p.slides.add_slide(self.p.slide_layouts[6])
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = WHITE
        self.n += 1
        if kicker:
            self.text(s, kicker.upper(), 0.6, 0.3, 6.5, 0.3, size=11, color=ACCENT, bold=True)
        if section:
            x = 12.75
            for name in reversed(self.SECTIONS):
                w = 0.3 + 0.085 * len(name)
                x -= w
                self.text(s, name, x, 0.3, w, 0.3, size=9, color=ACCENT if name == section else MUTED, bold=name == section,
                          align=PP_ALIGN.RIGHT)
        self.text(s, title, 0.6, 0.62, 12.2, 0.9, size=26 if len(title) <= 62 else 22 if len(title) <= 76 else 20, color=INK, bold=True)
        rule = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(1.32), Inches(0.9), Inches(0.05))
        rule.fill.solid(); rule.fill.fore_color.rgb = ACCENT; rule.line.fill.background(); rule.shadow.inherit = False
        self.text(s, f"{self.n}", 12.4, 7.08, 0.5, 0.3, size=9, color=MUTED, align=PP_ALIGN.RIGHT)
        self.text(s, "Did That Trade Really Happen?  ·  PW26_SSS_08  ·  Capstone Review 2", 0.6, 7.08, 8, 0.3, size=9, color=MUTED)
        s.notes_slide.notes_text_frame.text = notes
        return s

    def text(self, s, txt, x, y, w, h, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT, bullets=None):
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        lines = bullets if bullets is not None else [txt]
        for i, line in enumerate(lines):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.alignment = align
            if isinstance(line, tuple):
                b, rest = line
                r1 = para.add_run(); r1.text = b; self._fmt(r1, size, color, True)
                r2 = para.add_run(); r2.text = rest; self._fmt(r2, size, color, False)
            else:
                r = para.add_run(); r.text = (("•  " if bullets is not None else "") + line); self._fmt(r, size, color, bold)
            para.space_after = Pt(8 if bullets is not None else 0)
        return tb

    @staticmethod
    def _fmt(r, size, color, bold):
        r.font.size, r.font.color.rgb, r.font.bold, r.font.name = Pt(size), color, bold, FONT

    def box(self, s, txt, x, y, w, h, fill=WHITE, color=INK, size=15, bold=False, line=LINE, sub=None):
        sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        sh.adjustments[0] = 0.06
        sh.shadow.inherit = False
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
        sh.line.color.rgb = line; sh.line.width = Pt(1.25)
        tf = sh.text_frame; tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.1)
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = txt; self._fmt(r, size, color, bold or True)
        if sub:
            p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
            r2 = p2.add_run(); r2.text = sub; self._fmt(r2, size - 3, INK2, False)
        return sh

    def arrow(self, s, x1, y1, x2, y2):
        c = s.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        c.line.color.rgb = INK2; c.line.width = Pt(1.75)
        ln = c.line._get_or_add_ln()
        from pptx.oxml.ns import qn
        tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
        ln.append(tail)

    def image(self, s, name, x, y, w=None, h=None, crop_top=0.0, crop_bottom=0.0):
        path = SHOTS / name if (SHOTS / name).exists() else Path(name)
        pic = s.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w) if w else None, Inches(h) if h else None)
        if crop_top or crop_bottom:
            pic.crop_top, pic.crop_bottom = crop_top, crop_bottom
        pic.line.color.rgb = LINE
        return pic

    def fit(self, s, path, x, y, maxw, maxh, center=True):
        """place an image inside the box (x, y, maxw, maxh) inches, preserving aspect ratio - never overflows the slide."""
        from PIL import Image
        path = Path(path) if Path(path).is_absolute() else ROOT / "assets/ppt" / path
        w, h = Image.open(path).size
        scale = min(maxw / w, maxh / h)
        W, H = w * scale, h * scale
        ox = x + ((maxw - W) / 2 if center else 0)
        pic = s.shapes.add_picture(str(path), Inches(ox), Inches(y), Inches(W), Inches(H))
        pic.line.color.rgb = LINE
        return pic

    def verdict_chip(self, s, word, x, y, color, fill):
        sh = self.box(s, word, x, y, 1.35, 0.5, fill=fill, color=color, size=16, line=fill)
        return sh


def table(d, s, rows, x, y, w, col_w, size=13, row_h=0.42, header_fill=NAVY):
    """rows[0] = header. col_w = relative widths."""
    shp = s.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(row_h * len(rows)))
    tb = shp.table
    tot = sum(col_w)
    for j, cw in enumerate(col_w):
        tb.columns[j].width = Emu(int(Inches(w) * cw / tot))
    for i, row in enumerate(rows):
        tb.rows[i].height = Inches(row_h)
        for j, val in enumerate(row):
            c = tb.cell(i, j)
            c.fill.solid()
            c.fill.fore_color.rgb = header_fill if i == 0 else (WHITE if i % 2 else RGBColor(0xF4, 0xF6, 0xF8))
            c.margin_left = c.margin_right = Inches(0.06)
            c.margin_top = c.margin_bottom = Inches(0.03)
            tf = c.text_frame
            tf.word_wrap = True
            tf.paragraphs[0].text = ""
            r = tf.paragraphs[0].add_run()
            r.text = str(val)
            d._fmt(r, size, WHITE if i == 0 else INK, i == 0)
    return tb


def mono(d, s, lines, x, y, w, h, size=13):
    """monospace panel for formulas and pseudocode; comment lines (▷) in navy."""
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor(0xF7, 0xF8, 0xFA); sh.line.color.rgb = LINE; sh.shadow.inherit = False
    tf = sh.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.15)
    for i, ln in enumerate(lines):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = PP_ALIGN.LEFT
        r = para.add_run(); r.text = ln
        r.font.size, r.font.name, r.font.color.rgb = Pt(size), "Consolas", (NAVY if not ln.startswith(" ") and ln else INK)
        r.font.bold = bool(ln) and not ln.startswith(" ")
    return sh


PLAIN = RGBColor(0xFF, 0xF6, 0xDB)       # "in plain words" callouts


def plain(d, s, txt, x, y, w, h, size=15):
    """quiet 'in plain words' line: amber bar + text (the everyday version a non-specialist presenter can read out)."""
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y + 0.05), Inches(0.06), Inches(h - 0.1))
    bar.fill.solid(); bar.fill.fore_color.rgb = WARN; bar.line.fill.background(); bar.shadow.inherit = False
    tb = s.shapes.add_textbox(Inches(x + 0.15), Inches(y), Inches(w - 0.15), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]
    r1 = p.add_run(); r1.text = "In plain words:  "; d._fmt(r1, size, WARN, True)
    r2 = p.add_run(); r2.text = txt; d._fmt(r2, size, INK2, False)
    return tb


def N(say, *qa):
    """speaker notes = a word-for-word script + likely panel questions with simple answers."""
    out = "SAY:\n" + say
    if qa:
        out += "\n\nIF THE PANEL ASKS:"
        for q, a in qa:
            out += f"\nQ: {q}\nA: {a}"
    return out


GLOSSARY = [
    ("FACT + the 3 T's", "memory aid for the algorithm: Frames → clock, All claims, Check (Trade, Timing, Tenure), Tell the verdict. "
                         "Hook: place it in time, list what it claims, then ask - did it trade, was it there, is it old enough?"),
    ("Market record (tape)", "public 1-minute data for NSE: lowest and highest price (or index level) in each minute, and shares traded for "
                             "stocks. We use Yahoo Finance's copy - a data vendor, not NSE's own trade-by-trade record. It has no names."),
    ("Index level", "an index such as the Nifty 50 is a calculated number, not something bought and sold. We check that a stated level "
                    "was within the recorded range - not that a trade happened."),
    ("Knowability (price-timing consistency)", "was the stated price inside the market's recorded range just before it was said? It "
                                               "does NOT show what the speaker saw or knew."),
    ("Claim", "any price, quantity or trade mentioned in the video, e.g. 'Nifty is at 22,700'."),
    ("Encoder", "a part of the software that 'reads' one kind of input: the chart, the text on screen, or the speech."),
    ("OCR", "software that reads text from a picture. ASR: software that turns speech into text."),
    ("Ratchet / day-high", "the chart's 'highest price so far today'. It can only go up, like a high-water mark, so it tells us the time."),
    ("τ (tau) / time map", "how many seconds the video is behind or ahead of the real market clock."),
    ("PASS / FAIL / ABSTAIN", "fits the records / impossible under every sensible reading / not enough evidence to tell."),
    ("Assumption", "something we take as given, written down openly (e.g. 'the trade was on NSE')."),
    ("Baseline", "a simpler method we compare against - here mostly our own first version."),
    ("Ablation", "switching one part off to prove it is needed."),
    ("Held-out / in-sample", "tested on data never seen while building / tested on the same data used to fix it (weaker)."),
    ("Certificate", "the downloadable proof sheet: answers, evidence, assumptions and data fingerprints (hashes)."),
]

def wilson(k, n, z=1.96):
    """95 % Wilson score interval for k successes in n trials, in percent."""
    if n == 0:
        return (float("nan"), float("nan"))
    p = k / n
    den = 1 + z * z / n
    c = (p + z * z / (2 * n)) / den
    h = z * ((p * (1 - p) / n + z * z / (4 * n * n)) ** 0.5) / den
    return (100 * max(0.0, c - h), 100 * min(1.0, c + h))


def ci(k, n):
    lo, hi = wilson(k, n)
    f = lambda v: f"{v:.1f}" if 0 < v < 1 else f"{v:.0f}"
    return f"{k}/{n} = {100 * k / n:.0f} %  [{f(lo)}, {f(hi)}]"


DATA_LINE = "Reference data: NSE prices as 1-minute bars from Yahoo Finance (a data vendor) - not NSE's own trade-by-trade record"
LIGHT = RGBColor(0xEE, 0xF3, 0xF8)
PINK = RGBColor(0xF8, 0xE4, 0xE4)
MINT = RGBColor(0xE3, 0xF1, 0xEA)
FULL, HALF, NONE = "yes", "partly", "–"


def main():
    ex, e5 = load_numbers()
    abl = {r["recording"]: r for r in json.loads((ROOT / "reports/ablation_ratchet.json").read_text(encoding="utf-8"))["rows"]}
    e2, e3, e4 = ex["E2_spoken_heldout"], ex["E3_controlled_manipulations"], ex["E4_resources"]
    v1_con = sum(e2["v1_frozen_heldout"]["vs_labels"]["contradicted"].values())   # 37
    v1_ok = 1                                         # per-item inspection: only the 24,616.6 item (R28/R30)
    g = lambda k, v="Contradicted by available records": e3[k].get(v, 0)
    exact = g("B1 [exact] +2%") + g("B2 [exact] -2%") + g("B3 [exact] +5%")
    ineq = g("B1 [inequality] 2% in the falsifying direction") + g("B3 [inequality] 5% in the falsifying direction")
    bank = abl["sample:tv18_nifty_bank_1301"]["price_only"]
    lev = abl["sample:tv18_nifty_levels_1250"]["price_only"]
    d = Deck()

    # 1 ── title
    s = d.slide("Did That Trade Really Happen?", "UE23CS441A · Capstone Phase 3 · Review 2", notes=N(
        "Good morning. We are Adishree and Akshat, guided by Dr. Shylaja S S. When a market video states a price or a trade, is that possible "
        "according to the market's own record at that moment? The picture is the whole talk: read the time from the TV chart, look up the "
        "recorded minute, answer fits / impossible / can't tell, with the evidence.",
        ("Is this AI?", "AI models transcribe speech and read the screen; the checking itself is explicit, re-checkable logic.")))
    d.text(s, "Checking prices and trades stated in market videos against the market's own record, at that moment",
           0.6, 1.5, 12.2, 0.5, size=17, color=INK2)
    d.text(s, "① a real TV frame: the software reads the boxed numbers", 0.6, 2.05, 5.6, 0.3, size=11, color=ACCENT, bold=True)
    d.fit(s, "fig_frame_annotated.png", 0.6, 2.35, 5.6, 2.5, center=False)
    d.arrow(s, 6.3, 3.55, 6.75, 3.55)
    d.text(s, "② when was this?", 6.8, 2.05, 1.9, 0.3, size=11, color=ACCENT, bold=True)
    d.box(s, "6 Oct 2026\n13:02:55", 6.8, 3.0, 1.75, 0.75, size=14, fill=LIGHT, color=ACCENT, line=LIGHT)
    d.text(s, "give or take 4 seconds - worked out from the chart alone (how: slide 7). No chart on screen → our recording clock, labelled", 6.8, 3.8, 1.8, 1.0, size=10, color=MUTED)
    d.arrow(s, 8.6, 3.55, 9.0, 3.55)
    d.text(s, "③ what did the market record then?", 9.05, 2.05, 3.9, 0.3, size=11, color=ACCENT, bold=True)
    d.fit(s, "fig_tape_example.png", 9.05, 2.35, 3.9, 1.95, center=False)
    d.text(s, "bars = each minute's low–high · dashed = the 22,700 said · pink = its rounding range · blue = minutes the speaker could quote", 9.05, 4.3, 3.9, 0.45, size=9, color=MUTED)
    for i, (w, m, col, fill) in enumerate([("PASS", "fits the record", OK, MINT), ("FAIL", "impossible", BAD, PINK),
                                           ("ABSTAIN", "cannot tell", WARN, RGBColor(0xFB, 0xEF, 0xD6))]):
        d.box(s, w, 9.05 + i * 1.32, 4.85, 1.22, 0.4, size=12, fill=fill, color=col, line=fill)
        d.text(s, m, 9.05 + i * 1.32, 5.25, 1.22, 0.3, size=10, color=MUTED, align=PP_ALIGN.CENTER)
    d.text(s, "④ the answer: TIMING PASS (the other two checks, slide 3: can't tell)", 6.8, 4.75, 2.2, 0.8, size=10, color=ACCENT, bold=True)
    d.text(s, "", 0.6, 5.0, 6.0, 0.9, size=14, bullets=[
        ("Team: ", "Adishree Gupta (PES1UG23CS024) · Akshat (PES1UG23CS048)"),
        ("Guide: ", "Dr. Shylaja S S  ·  PES University  ·  PW26_SSS_08")])
    d.text(s, "Scope: on real TV we tested stated PRICES (‘Nifty is at 22,700’). Personal TRADE claims (‘I bought 500 shares’) do not occur on "
              "news TV, so those were tested only on 7 made-up examples. Market record = NSE prices per minute, via the data vendor Yahoo Finance. "
              "Unfamiliar word? See the glossary (last slide).", 0.6, 5.85, 12.3, 0.65, size=11, color=MUTED)
    plain(d, s, "we read the time from the TV chart itself, look up what the market recorded at that moment, and say fits / impossible / can't tell.",
          0.6, 6.55, 12.3, 0.45, size=13)

    # 2 ── motivation
    s = d.slide("A wrong number on air looks exactly like a right one", "Real-world motivation", section="Motivation", notes=N(
        "Two real cases. First, our own CNBC-TV18 recording: 'Bank Nifty now at 55,150' - the recorded minute was 55,114 to 55,157, so it fits. "
        "Second, NDTV Profit's YouTube captions: 'Nifty50 24,616.6' - the record says 22,616.6. One digit differs, so it is most likely a caption "
        "error or a slip of the tongue; either way the number viewers read is impossible. And automatic checking is risky: our own first "
        "version wrongly called 36 of 37 statements impossible.",
        ("So did the anchor make a mistake?", "We cannot say. The caption text is impossible; who caused it the record cannot show.")))
    table(d, s, [["Said or shown on TV", "Where and when", "What the market recorded", "Our answer"],
                 ["“Bank Nifty now at 55,150”", "CNBC-TV18, our own recording, 6 Oct 2026, about 13:03", "13:03 minute: 55,114.55 to 55,156.90", "PASS - it fits"],
                 ["“Nifty50 24,616.6” (YouTube's automatic captions)", "NDTV Profit live stream on YouTube, 30 Sep 2026, 15:17:43",
                  "15:16–15:19: 22,616.6 (whole day: 22,595–22,809)", "FAIL - impossible; one digit off, so a caption error or slip of the tongue"]],
          0.6, 1.65, 12.2, [2.8, 3.5, 3.3, 2.6], size=12, row_h=0.72)
    for i, (h, b) in enumerate([("Viewers act on these numbers", "checking by hand needs the exact minute, the exact stock or index, and how the speaker rounded"),
                                ("Even our own first version got it wrong", "of 37 statements it called ‘impossible’, 36 were not - careless checking accuses honest speakers (slide 11)"),
                                ("Regulation now depends on it", "SEBI rule from 1 Jul 2026: someone who only teaches about stocks must not use prices from the last 30 days when recommending one - so the age of prices shown in a video matters")]):
        d.box(s, h, 0.6 + i * 4.1, 4.1, 3.9, 0.6, size=13, fill=ACCENT, color=WHITE, line=ACCENT)
        d.box(s, b, 0.6 + i * 4.1, 4.75, 3.9, 1.2, size=12, fill=WHITE)
    plain(d, s, "a commentator says ‘Kohli is on 87’ - you check the scoreboard. Markets keep a scoreboard too; we check videos against it.",
          0.6, 6.3, 12.2, 0.55, size=13)

    # 3 ── problem
    s = d.slide("The task: for every claim, is it possible, impossible, or can't we tell?", "Problem statement", section="Problem", notes=N(
        "Given a video, the market's minute record and the date the video was published, we give every claim three answers - Trade, Timing, "
        "Tenure - with evidence. The key rule: say FAIL only when the claim is impossible however the sentence is reasonably read, and only "
        "when all four conditions on the right hold. Otherwise ABSTAIN. Note the app's names for the three checks, shown in brackets.",
        ("Who decides what a 'sensible reading' is?", "The software tries the readings a listener could reasonably have: which stock, whether "
                                                      "the number is a price now or a target, and how it was rounded. All must fail for FAIL.")))
    d.box(s, "In", 0.6, 1.6, 0.9, 0.55, size=14, fill=ACCENT, color=WHITE, line=ACCENT)
    d.box(s, "the video  +  the market's minute-by-minute record  +  the publishing date (used only by TENURE)", 1.6, 1.6, 6.3, 0.55, size=12, fill=WHITE)
    d.box(s, "Out", 0.6, 2.25, 0.9, 0.55, size=14, fill=ACCENT, color=WHITE, line=ACCENT)
    d.box(s, "for every claim: three answers (PASS / FAIL / ABSTAIN) + the evidence behind each", 1.6, 2.25, 6.3, 0.55, size=12, fill=WHITE)
    d.box(s, "Rule", 0.6, 2.9, 0.9, 0.55, size=14, fill=BAD, color=WHITE, line=BAD)
    d.box(s, "say FAIL only if the claim is impossible however it is sensibly read", 1.6, 2.9, 6.3, 0.55, size=12, fill=WHITE)
    d.text(s, "The three checks - ‘the 3 T's’ (app names in brackets, used on slide 13)", 0.6, 3.6, 7.3, 0.35, size=13, color=ACCENT, bold=True)
    for i, (t, app, q) in enumerate([("TRADE", "Fill conservation", "could that many shares have changed hands at that price, that day, on NSE?"),
                                     ("TIMING", "Knowability", "was the stated price within what the market recorded just before it was said?"),
                                     ("TENURE", "Data age", "(= data age) was the price data at least 30 days old? SEBI wants teachers to use OLD prices, so lessons are not live tips")]):
        d.box(s, t, 0.6, 4.0 + i * 0.66, 1.5, 0.56, size=13, fill=ACCENT, color=WHITE, line=ACCENT, sub=None)
        d.box(s, f"{q}  (app: ‘{app}’)", 2.2, 4.0 + i * 0.66, 5.7, 0.56, size=11, fill=WHITE)
    d.box(s, "A price claim gets FAIL only if ALL four hold", 8.2, 1.6, 4.7, 0.5, fill=BAD, color=WHITE, size=13, line=BAD)
    d.text(s, "", 8.2, 2.15, 4.7, 3.7, size=12, bullets=[
        ("1  Complete data: ", "at least 90 % of the minutes checked have records (gaps could hide the price)"),
        ("2  Clear claim: ", "the stock or index is named in the same sentence and the number is stated as its price NOW - not a target, "
                             "a support level, an old purchase, a prediction, or a number the speaker corrected (‘23,000 - sorry, 22,672’)"),
        ("3  Every reading fails: ", "the number is outside what was recorded, whichever stock, rounding or meaning a listener could take"),
        ("4  Bigger than data noise: ", "data providers report a day's high or low slightly differently; differences under 0.25 % are treated as noise"),
        ("Otherwise → ABSTAIN.", " TRADE also fails if the market was closed or more shares are claimed than traded; TENURE never fails.")])
    d.box(s, "Each check is answered separately (any FAIL = the claim is impossible). Never decided: who traded · whether the video is genuine · whether anyone broke the law", 0.6, 6.05, 12.3, 0.42, size=12,
          fill=PINK, color=BAD, line=PINK)
    plain(d, s, "we only say ‘impossible’ when every sensible way of reading the sentence fails.", 0.6, 6.52, 12.3, 0.45, size=13)

    # 4 ── related work
    s = d.slide("No earlier method both dates a video from its own screen and checks it carefully",
                "Research gap and existing work", section="Problem", notes=N(
        "Each row is a family of published research methods; the names are the papers' names, listed on slide 15. Read the table by column. "
        "The closest idea to ours is ENF forensics: it works out when an audio recording was made by matching the faint electric hum in it to "
        "the power grid's log. We do the same kind of thing with the price and day's high on a TV chart, matched to the market's log.",
        ("Isn't this just ENF?", "Same family: dating against a reference log. Different signal - the TV chart - plus careful claim checking.")))
    cols = ["Facts that change minute to minute", "Works out WHEN the video was made", "Can say ‘cannot tell’", "Uses stock-market data", "Works on video"]
    rows = [("Fact-checkers [1–3] - check a sentence against Wikipedia or a table", [NONE, NONE, HALF, NONE, NONE]),
            ("Number fact-checker [4] - checks numbers in news text", [HALF, NONE, NONE, NONE, NONE]),
            ("Fake-photo detector [5] - spots a photo used with the wrong caption", [NONE, NONE, NONE, NONE, HALF]),
            ("Electric-hum dating [9] - dates audio from the grid hum in it (closest idea)", [FULL, FULL, HALF, NONE, FULL]),
            ("Chart readers [10, 11] - turn a chart image into numbers", [NONE, NONE, NONE, HALF, NONE]),
            ("Time-series matching [12] - lines up two series, one best guess", [FULL, HALF, NONE, NONE, NONE]),
            ("‘When to refuse to answer’ theory [6, 7]", [NONE, NONE, FULL, NONE, NONE]),
            ("This project (designed for all five)", [FULL, FULL, FULL, FULL, FULL])]
    tb = table(d, s, [["Earlier research (numbers = references, slide 15)"] + cols] + [[r] + v for r, v in rows], 0.6, 1.6, 12.2,
               [4.7] + [1.5] * 5, size=10.5, row_h=0.41)
    for j in range(1, 6):
        for i in range(len(rows) + 1):
            tb.cell(i, j).text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    for j in range(6):
        c = tb.cell(len(rows), j); c.fill.solid(); c.fill.fore_color.rgb = LIGHT
        for r in c.text_frame.paragraphs[0].runs:
            r.font.bold = True; r.font.color.rgb = ACCENT
    d.text(s, "Electric-hum dating: mains power hums at about 50 Hz with tiny drifts that the grid logs; a recording picks up the hum, so matching its drifts to the log reveals when it was made. How to read: ‘yes’ = the method does what the column says · ‘partly’ = in a limited way · ‘–’ = no. Our own reading of the papers; "
              "the last row is what this project was designed to do, not a claim that it is better at each.", 0.6, 5.75, 12.2, 0.6, size=10, color=MUTED)
    d.text(s, "Where we stand: we borrow the electric-hum idea - date a recording by matching it to a log - and apply it to market TV using the "
              "chart's price and day's high, plus a careful rule for saying ‘impossible’.", 0.6, 6.45, 12.2, 0.55, size=12, color=ACCENT, bold=True)

    # 5 ── contributions
    s = d.slide("Our contributions: a new question, dating TV from its chart, real evidence",
                "Research objectives and contributions", section="Problem", notes=N(
        "Three contributions, each tied to evidence. One: the question itself, with the rule that 'impossible' must survive every sensible "
        "reading. Two: dating from the chart. Three: measurements on real broadcasts. The software is engineering, not claimed as research novelty.",
        ("Which result would you defend first?", "That a careless checker was right in only 1 of 37 'impossible' calls on new TV days.")))
    for i, (t, body, tech, ev) in enumerate([
            ("1 · A new checking task", "For every claim in a video: PASS, FAIL or ABSTAIN against the market record - and FAIL only if impossible "
             "under every sensible reading.", "technical name: sound three-valued verification",
             "Evidence: guarantee ‘never wrongly impossible’ (slide 7); fair test (slide 11)"),
            ("2 · Dating TV from its own chart", "The chart's ‘highest so far today’ can only rise. Matching it and the price, frame by frame, "
             "with the record fixes the date and time - or the software says it cannot.", "technical name: monotone-display dating",
             "Evidence: guarantee ‘true time kept’ (slide 7); slides 11–12"),
            ("3 · Measured on real TV", "A fair test of our first version on TV days it never saw, deliberately changed quotes, and dates checked against our own "
             "recording clock - each with how many cases and how uncertain the result is.", "technical name: empirical study",
             "Evidence: slides 10–12")]):
        x = 0.6 + i * 4.1
        d.box(s, t, x, 1.65, 3.9, 0.6, fill=ACCENT, color=WHITE, size=15, line=ACCENT)
        d.box(s, body, x, 2.3, 3.9, 2.1, size=14, fill=WHITE)
        d.text(s, "for specialists - " + tech.replace("technical name: ", ""), x, 4.45, 3.9, 0.35, size=10, color=MUTED, align=PP_ALIGN.CENTER)
        d.box(s, ev, x, 4.85, 3.9, 0.7, size=12, fill=LIGHT, color=ACCENT, line=LIGHT)
    d.text(s, "Also built (engineering, not claimed as research): the working software - readers, checks, downloadable evidence sheets, "
              "a desktop interface and 27 automatic self-tests of the code.", 0.6, 5.75, 12.3, 0.5, size=12, color=MUTED)
    plain(d, s, "a new question, a way to read the clock off the TV chart, and honest measurements on real TV.", 0.6, 6.45, 12.3, 0.45, size=13)

    # 6 ── architecture
    s = d.slide("How it works: three readers build a list of claims; three checks judge it",
                "System architecture", section="Method", notes=N(
        "Pictures go to two readers: the frame reader reads the chart and works out the market time; the text reader reads prices, quantities "
        "and profit figures. Sound goes to the audio reader. Everything becomes one claim list with market times. The three checks run "
        "independently against the market record, and their answers form the certificate.",
        ("What happens without a chart on screen?", "The frame reader says so, and claims are timed with the recording computer's clock, labelled.")))
    d.box(s, "Video pictures", 0.35, 1.95, 1.6, 0.85, size=13, fill=LIGHT, line=LIGHT)
    d.box(s, "Video sound", 0.35, 4.6, 1.6, 0.85, size=13, fill=LIGHT, line=LIGHT)
    for i, (t, sub) in enumerate([("Frame reader", "reads the chart → date and time"), ("Text reader", "on-screen prices, share counts, profit figures"),
                                  ("Audio reader", "speech → words with times")]):
        d.box(s, t, 2.4, 1.65 + i * 1.45, 2.5, 1.1, size=14, sub=sub)
        d.arrow(s, 4.9, 2.2 + i * 1.45, 5.45, 3.6)
    d.arrow(s, 1.95, 2.37, 2.4, 2.2); d.arrow(s, 1.95, 2.37, 2.4, 3.65); d.arrow(s, 1.95, 5.02, 2.4, 5.1)
    d.box(s, "Claim list", 5.45, 3.05, 2.3, 1.1, size=15, sub="who said what, which stock, when", line=ACCENT)
    d.box(s, "Market record", 5.45, 5.0, 2.3, 1.0, size=14, sub="NSE, minute by minute (Yahoo Finance) → used by all three checks")
    d.arrow(s, 7.75, 5.5, 8.3, 5.1)
    for i, (t, sub) in enumerate([("TRADE", "enough shares traded?"), ("TIMING", "price recorded just before?"), ("TENURE", "how old is the data?")]):
        d.box(s, t, 8.3, 1.65 + i * 1.45, 2.4, 1.1, size=14, sub=sub)
        d.arrow(s, 7.75, 3.6, 8.3, 2.2 + i * 1.45)
        d.arrow(s, 10.7, 2.2 + i * 1.45, 11.1, 3.6)
    d.box(s, "Certificate", 11.1, 3.0, 1.85, 1.2, size=14, sub="the 3 answers + evidence, downloadable", line=ACCENT)
    d.text(s, "Terms: a ‘reader’ (technical word: encoder) turns one kind of input into facts · text in pictures is read by OCR (optical character "
              "recognition) · speech is turned into words by ASR (automatic speech recognition).", 0.6, 6.05, 12.3, 0.45, size=10, color=MUTED)
    plain(d, s, "two readers watch the screen, one listens; they write one shared list; three checks judge it and sign one report.", 0.6, 6.5, 12.3, 0.45, size=13)

    # 7 ── formulation + guarantees, in words (symbols in Appendix A)
    s = d.slide("Three rules and two guarantees - in words", "Proposed mathematical formulation", section="Method", notes=N(
        "The rules in words; the same rules in mathematical notation are in Appendix A for specialists. Rule 1 keeps only the moments where "
        "every frame's price and day-high match the record. Rule 2 says when a claim fits. Rule 3 gives the answer. The two guarantees are "
        "proven in Appendix B. The figure shows a single real frame: the price alone fits 22 minutes of the day; with the day's high, only 8.",
        ("What is the objective function?", "There is nothing to train. The rules are built so that FAIL is never issued wrongly under the stated conditions.")))
    for i, (h, b) in enumerate([("Rule 1 · find the time", "Keep only the moments - on any trading day - where the price AND the day's high "
                                 "shown in EVERY frame match what the market recorded."),
                                ("Rule 2 · does a claim fit?", "Yes if, allowing for rounding, the stated price was within the recorded range in "
                                 "a minute the speaker could have been looking at (up to 3 minutes before speaking)."),
                                ("Rule 3 · the answer", "FAIL if no sensible reading fits and the 4 conditions of slide 3 hold · PASS if at "
                                 "least one reading fits · otherwise ABSTAIN.")]):
        y = 1.6 + i * 0.95
        d.text(s, h, 0.6, y, 6.6, 0.3, size=13, color=ACCENT, bold=True)
        d.text(s, b, 0.6, y + 0.3, 6.6, 0.6, size=12)
    d.box(s, "Guarantee ‘true time kept’", 0.6, 4.5, 6.6, 0.38, size=12, fill=ACCENT, color=WHITE, line=ACCENT)
    d.text(s, "The real moment always survives Rule 1. So a wrong date can only come from a misread screen or a TV showing a different "
              "market - not from the rules themselves.", 0.6, 4.9, 6.6, 0.55, size=11)
    d.box(s, "Guarantee ‘never wrongly impossible’", 0.6, 5.45, 6.6, 0.38, size=12, fill=ACCENT, color=WHITE, line=ACCENT)
    d.text(s, "FAIL needs every reading we try to fail. So a wrong FAIL can only happen if the right reading was not among those tried, or "
              "the market data were wrong.", 0.6, 5.85, 6.6, 0.45, size=11)
    d.box(s, "Not guaranteed: a PASS only means ‘fits the record’ - it does not prove the claim true.", 0.6, 6.35, 6.6, 0.5, size=12,
          fill=PLAIN, color=INK, line=WARN)
    d.text(s, "What one real frame tells us (6 Oct 2026)", 7.45, 1.6, 5.5, 0.3, size=12, color=ACCENT, bold=True)
    d.fit(s, "fig_insight_frame.png", 7.45, 1.9, 5.45, 2.4, center=False)
    d.text(s, "How to read: grey bars = lowest–highest index level in each minute of the day · brown line = the day's high so far (it only "
              "rises) · red ticks = minutes where the shown price alone fits · green bands = minutes where price AND day's high both fit. "
              "One frame narrows the day from 22 possible minutes to 8; each further frame narrows it again. The arrow marks the final answer after all 56 frames.",
           7.45, 4.35, 5.45, 1.7, size=11, color=INK2)
    d.text(s, "Formulas: Appendix A · proofs: Appendix B (the last slides, for specialists).", 7.45, 6.1, 5.45, 0.3, size=10, color=MUTED)

    # 8 ── algorithm: FACT steps, each in plain words AND pseudocode
    s = d.slide("The algorithm: four steps (FACT), in words and in pseudocode", "Proposed algorithm", section="Method", notes=N(
        "Remember FACT. Frames to clock: the market time of every frame, from the chart. All claims: what the video shows and says. Check with "
        "three T's: Trade, Timing, Tenure. Tell the verdict with evidence. Middle column: what each step does in plain words. Right column: "
        "the same step as pseudocode - exactly what the software runs. The three distinctive operations are named at the bottom.",
        ("Why have a Tenure check that never FAILs?", "It reports a fact - how old the data was. Whether the SEBI rule applies depends on who "
                                                      "the speaker is, which a video cannot show; so it never declares a breach."),
        ("What is the cost?", "Per trading day, each frame's set is built in one pass over the day's minutes; intersecting sets is linear in "
                              "their size. All checks run in under a second on a laptop.")))
    d.text(s, "In plain words", 2.85, 1.5, 4.4, 0.3, size=12, color=ACCENT, bold=True)
    d.text(s, "Pseudocode (what the software runs)", 7.35, 1.5, 5.6, 0.3, size=12, color=ACCENT, bold=True)
    steps = [("F", "Frames → clock", 1.45,
              "Read the price and the day's high from the chart in every frame. For every trading day, keep only the moments that match ALL "
              "frames. Exactly one day and within ±90 s → use it (typically ±4 s at 1 picture per second). Otherwise say why and use our "
              "recording clock, clearly labelled.",
              [" for d in trading_days:",
               "   T[d] ← whole session",
               "   for each frame f:   p, h ← OCR(chart price, day-high label)",
               "     T[d] ← T[d] ∩ {τ : p ∈ bar(t_f+τ) ± 0.10",
               "                       and RunMax(t_f+τ) ∈ [h−1, h+1]}",
               " if exactly one d has T[d] ≠ ∅ and width(T[d]) ≤ ±90 s:",
               "   τ ← T[d]          (IDENTIFIED)",
               " else: τ ← capture clock ± u   (LABELLED, reason shown)"]),
             ("A", "All claims", 0.95,
              "Turn speech into words and read on-screen text. For every number: which stock or index, the value, how many shares, the "
              "rounding, and what kind - a price now, an old price, a target, a recommendation, a level or a self-correction.",
              [" words ← ASR(audio);  text ← OCR(frames)",
               " for each number n in words ∪ text:",
               "   c ← (instrument, value P, shares q, rounding δ, kind)",
               "   kind ∈ {now, past, target, call, level, retracted}",
               "   C ← C ∪ {c}"]),
             ("C", "Check: the 3 T's", 1.45,
              "TRADE: impossible if the market was closed, the price never traded, or more shares are claimed than all of NSE traded then. "
              "TIMING: fits if the price was recorded in the minutes before it was said; FAIL only if every reading fails (4 conditions, "
              "slide 3). TENURE: reports whether the data was 30+ days old; never FAIL.",
              [" for c in C:   m ← t_c + τ",
               "   TRADE : FAIL if holiday(m) or no bar near m holds P",
               "           or q > Σ volume(bars holding P)    (NSE)",
               "   TIMING: ok(a) ← some bar in window W_a(m) holds P ± δ",
               "           FAIL if complete ∧ clear(c) ∧ ∀a: ¬ok(a)",
               "           PASS if ∃a: ok(a);  else ABSTAIN",
               "   TENURE: PASS if age ≥ 30 d;  else ABSTAIN (never FAIL)"]),
             ("T", "Tell the verdict", 0.72,
              "Write the certificate: every answer with the exact market minutes, words, frames and assumptions, plus a fingerprint of each "
              "data file so anyone can re-check.",
              [" cert ← {verdicts, deciding bars, words, frames,",
               "         assumptions, sha256(data files)};   return cert"])]
    y = 1.85
    for L, t, h, b, code in steps:
        d.box(s, L, 0.6, y, 0.6, h - 0.08, size=22, fill=ACCENT, color=WHITE, line=ACCENT)
        d.box(s, t, 1.25, y, 1.5, h - 0.08, size=12, fill=WHITE)
        d.text(s, b, 2.85, y - 0.02, 4.4, h, size=10.5)
        mono(d, s, code, 7.35, y, 5.6, h - 0.08, size=9)
        y += h
    plain(d, s, "what is distinctive: (1) the rising day-high pins the clock; (2) ‘impossible’ needs EVERY reading to fail; (3) traded volume caps any claimed trade.", 0.6, 6.5, 12.3, 0.4, size=11)

    # 9 ── worked example
    s = d.slide("Worked example: one real claim from pixels to verdict", "Working example · CNBC-TV18 · 6 Oct 2026", section="Method", notes=N(
        "One real sentence end to end. The anchor says '22,700 is where the index finds itself'. The frame reader reads the boxed numbers; of "
        "the 7 trading days in our data only 6 October fits, at 13:02:55 give or take 4 seconds. The speaker may be reading a screen up to 3 "
        "minutes old, so we check 12:59:51 to 13:02:59. The 13:02 minute ranged 22,693 to 22,706; allowing for rounding, 22,700 fits: TIMING "
        "PASS. No share count was said, so TRADE can't tell; TENURE can't tell because the rule depends on who the speaker is.",
        ("Why 22,699.40 in the picture but 22,699.35 elsewhere?", "The on-screen number changes several times a second; the picture is a "
                                                                  "frame a fraction of a second apart. Both fit.")))
    for i, (a, b) in enumerate([("Said", "“22,700 is where the index finds itself” - a rounded number: anything from 22,650 to 22,750"),
                                ("Shown", "Nifty 50 at 22,699.40 (it changes every split second; a frame earlier: 22,699.35); day's high so far 22,729"),
                                ("When", "of the 7 trading days in our minute data (29 Sep–8 Oct) only 6 Oct fits → said at 13:02:55, give or take 4 s"),
                                ("Record", "check from 13:02:55 − 3 min − 4 s = 12:59:51 to 13:02:55 + 4 s = 13:02:59; the 13:02 minute alone (22,693.15–22,706.20) contains 22,700"),
                                ("Answers", "TIMING PASS · TRADE can't tell (no share count said) · TENURE can't tell (depends on who the speaker is) → no check fails")]):
        d.box(s, a, 0.6, 1.6 + i * 0.95, 1.35, 0.8, size=13, fill=ACCENT, color=WHITE, line=ACCENT)
        d.box(s, b, 2.05, 1.6 + i * 0.95, 4.75, 0.8, size=11, fill=WHITE if i < 4 else MINT)
    d.fit(s, "fig_frame_annotated.png", 7.0, 1.55, 5.95, 2.7)
    d.fit(s, "fig_tape_example.png", 7.0, 4.3, 5.95, 2.55)
    plain(d, s, "22,700 was within the range the market recorded at that minute - so it fits. That is all a PASS says.", 0.6, 6.45, 6.3, 0.5, size=12)

    # 10 ── setup
    s = d.slide("How we tested it: six experiments, each answering one question", "Experimental setup", section="Evidence", notes=N(
        "Each experiment answers one question with its own data, so we never quote one overall accuracy. E2 is the fair test: the first version "
        "was locked and scored on 8 TV days it had never seen. Fixes were then made using those same days, so results after the fixes are weaker evidence.",
        ("An AI judged an AI?", "A separate AI model was given only each sentence and asked what it claims - a current price, a target and so "
                                "on. It never saw our answers. Human judges are a next step.")))
    table(d, s, [["", "Question", "Data (how many)", "Kind of evidence"],
                 ["E1", "Can the scrolling ticker text be trusted?", "8 NDTV Profit recordings", "real"],
                 ["E2", "When the system says FAIL, is it right?", "8 TV days it never saw · all 105 of its firm answers judged (37 FAIL + 68 PASS). The first version also judged old prices, levels and targets (55 of the 105) - that was its mistake; the final version does not",
                  "fair test of the first version; then fixes made on those days (weaker)"],
                 ["E3", "Are deliberately changed numbers caught?", "50 real quotes the system had passed, each changed several ways: 29 exact quotes × 3 changes = 87; 21 ‘above/below’ quotes × 2 = 42; moved to the day before: 29 + 21", "made-up changes"],
                 ["E4", "How fast? How much memory?", "demo clips; 5.2 hours of audio", "measured on a laptop"],
                 ["E5", "Is the date right? Does the time match a separate clock?", "4 full 20-minute recordings, one picture every 10 seconds", "real"],
                 ["E6", "Does the chart's day-high matter?", "the same 4 recordings + 2 short clips", "real; run again with the day-high ignored"]],
          0.6, 1.6, 12.2, [0.6, 4.0, 5.0, 2.6], size=12, row_h=0.52)
    d.text(s, "", 0.6, 5.62, 12.2, 1.4, size=10.5, bullets=[
        ("Fair test (‘held-out’): ", "scored on data never used while building, with the system locked (‘frozen’). After fixing on the same data: weaker evidence."),
        ("Who judged right or wrong: ", "a separate AI model given only each sentence, never our answers - not human experts (a stated limitation)."),
        ("Compared against (‘baselines’ = simpler methods): ", "our locked first version (slide 11) · dating by price only (slide 12) · a simpler ticker rule (slide 12) · our recording computer's clock (slide 11)."),
        ("Equipment and data: ", "an ordinary laptop, no graphics card · NSE prices per minute from Yahoo Finance.")])

    # 11 ── results
    s = d.slide("Our own careless first version was almost always wrong when it said ‘impossible’", "Quantitative results", section="Evidence", notes=N(
        "Left: the fair test. The locked first version said FAIL 37 times on new days; only 1 was right. After the fixes, on the same days, one "
        "FAIL remains and it is right - weaker evidence, because those days were used for fixing. Right: each test with its count and a 95 "
        "percent range. Dating: right for all 3 recordings where it gave a date; it refused one.",
        ("Is the improvement proven?", "Not yet on new days. It is shown on the days used for fixing; a fresh locked test is the first next step.")))
    d.fit(s, "fig_fail_counts.png", 0.6, 1.55, 5.6, 2.4, center=False)
    d.text(s, "How to read: bar length = number of ‘impossible’ (FAIL) answers; green = right, red = wrong because the sentence was misunderstood.",
           0.6, 3.95, 5.6, 0.45, size=10, color=MUTED)
    d.text(s, "", 0.6, 4.4, 5.6, 2.0, size=12, bullets=[
        ("Fair test, first version: ", f"FAIL right {v1_ok} of {v1_con} times"),
        ("After the fixes (same days - weaker): ", "1 FAIL, and it is right · of 50 real ‘price now’ statements: 43 PASS · 1 FAIL · 6 ABSTAIN"),
        ("That 1 FAIL: ", "captions say ‘24,616.6’, the record says 22,616.6 - impossible as written")])
    table(d, s, [["Test", "Result  [95 % range]"],
                 ["Exact quotes changed by ±2 %, +5 % (made up)", ci(exact, 87) + " caught"],
                 ["‘Above/below X’ quotes pushed the wrong way (made up)", ci(ineq, 42) + " caught"],
                 ["Quote moved to the previous day - exact (made up)", ci(g("C1 [exact] previous trading day"), 29) + " caught"],
                 ["Quote moved to the previous day - above/below (made up)", ci(g("C1 [inequality] previous trading day"), 21) + " caught - weak: the same price often occurs the day before too"],
                 ["Unchanged quotes (made up)", "50 of 50 pass - expected, they were picked from quotes that passed; only shows nothing broke"],
                 ["Date right (real)", f"{e5['n_day_correct']} of {e5['n_identified']} recordings where it gave a date; for 1 of the {e5['n_recordings']} recordings it refused"],
                 ["Time vs our recording clock (real)", "agreed within 6.2 s · reading 1 picture every 10 s gives ±15–30 s; 1 picture every second gives ±4 s"],
                 ["Scrolling ticker (real)", "readable reliably in 1 of 8 recordings - so ticker text is ignored elsewhere"],
                 ["Speed (E4)", f"each check under 1 second; {e4['process_rss_MB']:.0f} MB memory; ordinary laptop"],
                 ["FINAL version on new days", "NOT YET MEASURED - no unused days left in our archive; first next step"]],
          6.45, 1.55, 6.45, [2.9, 3.55], size=10, row_h=0.42)
    d.text(s, "[95 % range] = with this few cases, the true rate very likely lies between the two numbers.", 6.45, 6.15, 6.45, 0.3, size=10, color=MUTED)
    plain(d, s, "the careless first version was wrong 36 of 37 times it said ‘impossible’. The final version's accuracy on new days is NOT yet measured.",
          0.6, 6.5, 12.3, 0.45, size=13)

    # 12 ── ablation
    s = d.slide("Without the chart's day-high, short clips cannot be dated (long recordings: no difference)", "Ablation and robustness", section="Evidence", notes=N(
        "We switched parts off to show they are needed. The curve: on a 110-second Bank Nifty clip, each frame rules out more possible times. With "
        "the day's high, one moment remains, give or take 4 seconds. With price alone, two separate moments about 24 minutes apart remain, so "
        "we cannot tell which. On 20-minute recordings it made no difference, which we report.",
        ("Why keep a part that does nothing on long recordings?", "Clips shared online are short; there it decides whether the time can be known.")))
    d.fit(s, "fig_where.png", 0.6, 1.6, 6.6, 2.3, center=False)
    d.text(s, "How to read: each bar is a moment the clip could be from. The clip is 110 seconds of CNBC-TV18 from about 13:01; its chart shows the Nifty 50."
              "  With the day's high there is one answer; with the price alone, two answers 24 minutes apart.",
           0.6, 4.0, 6.6, 0.75, size=11, color=INK2)
    d.text(s, "", 0.6, 4.8, 6.6, 1.3, size=11, bullets=[
        ("A second clip (12:50): ", f"price alone cannot even tell the DAY ({lev['candidate_days']} possible days); with the day's high, 1 day."),
        ("Long 20-minute recordings: ", "no difference - with many pictures the price alone is enough. So the day's high matters for short clips.")])
    d.text(s, "Switch one part off - what went wrong", 7.45, 1.55, 5.45, 0.3, size=12, color=ACCENT, bold=True)
    table(d, s, [["Part switched off", "What went wrong (measured)"],
                 ["Careful reading of sentences", f"‘impossible’ right only {v1_ok} of {v1_con} times"],
                 ["Testing whether a channel's scrolling ticker can be trusted", "≈ 12 % of ticker readings wrongly ‘impossible’"],
                 ["Requiring a number to be read the same in 2 pictures", "a one-off misreading (Titan read as 5,080; real price ~4,560) → wrongly ‘impossible’"],
                 ["0.25 % allowance between data providers", "a true ‘yesterday's high was 55,220’ (our data: 55,191) → wrongly ‘impossible’"],
                 ["Counting the number of shares", "an impossible 50-lakh-share claim would wrongly PASS"]],
          7.45, 1.9, 5.45, [2.7, 2.75], size=10.5, row_h=0.42)
    d.text(s, "Hard cases: it refuses rather than guesses", 7.45, 5.15, 5.4, 0.3, size=12, color=ACCENT, bold=True)
    d.text(s, "", 7.45, 5.45, 5.45, 1.0, size=10.5, bullets=[
        "chart matches no day → ABSTAIN · too few chart frames → time marked ‘weak’, recording clock used (labelled)",
        "wrong stock → ABSTAIN · targets, option prices (‘the 22,700 call’), self-corrections → not judged"])
    plain(d, s, "on a short clip, the day's high turns ‘12:38 or 13:03?’ into ‘13:03, give or take 4 seconds’.", 0.6, 6.45, 12.3, 0.45, size=13)

    # 13 ── UI
    s = d.slide("In the app, every answer is one click from its evidence", "Application UI and explainability", section="System", notes=N(
        "Real screenshots. Note the app's names for the three checks: Fill conservation is TRADE, Knowability is TIMING, Data age is TENURE. "
        "Top: the three answers. Middle: the TIMING answer in words. Bottom: the exact words spoken and the frame at that moment. The 40 percent is "
        "how sure the speech software was about the number; the number also appears on screen (22,699), so the claim is clear - but the app "
        "does not yet lower its answers for low confidence, which is listed as future work.",
        ("Why PASS with only 40 % confidence?", "The confidence is shown, not hidden; the screen independently shows 22,699. Using confidence to "
                                                 "switch to ABSTAIN is future work.")))
    d.fit(s, "certbar.png", 0.6, 1.55, 8.2, 0.6, center=False)
    d.box(s, "App names: Fill conservation = TRADE · Knowability = TIMING · Data age = TENURE", 8.95, 1.55, 3.95, 0.6, size=11, fill=LIGHT,
          color=ACCENT, line=LIGHT)
    d.fit(s, "trace_1_certificate.png", 0.6, 2.25, 12.3, 1.35, center=False)
    d.fit(s, "trace_3_encoders.png", 0.6, 3.7, 8.2, 2.95, center=False)
    d.text(s, "", 9.0, 3.7, 3.9, 3.2, size=11, bullets=[
        ("top bar: ", "the three answers; ‘Certificate e47b…’ = the evidence sheet's fingerprint"),
        ("box ‘1 · Certificate’ (box 2, the claim card, is not shown): ", "the TIMING answer in words - the minutes checked (up to 3 min = 180 s before speaking) and that the time came from the chart's day-high (the app calls it ‘ratchet’)"),
        ("box ‘3 · Encoders’, left: ", "the exact words; ‘confidence 40 %’ = how sure the speech software was of hearing ‘22,700’. Shown, not hidden - but not yet used to switch to ‘can't tell’ (a limitation, slide 14)"),
        ("box ‘3 · Encoders’, right: ", "the frame at that moment - it independently shows 22,699.35"),
        ("app words: ", "encoders = readers · ratchet = day's high · IST = Indian time · τ ±4 s = time give or take 4 s · 180 s = 3 min")])

    # 14 ── limitations
    s = d.slide("What this evidence does not show - and what comes next", "Limitations and future research", section="Outlook", notes=N(
        "Limits first, then the research agenda. The single biggest gap: the improved rules have not yet been tested on new days.",
        ("What is the first next experiment?", "Lock the current rules and score them on new TV days with human judges.")))
    d.text(s, "Not shown by this evidence", 0.6, 1.55, 6.1, 0.4, size=15, color=BAD, bold=True)
    d.text(s, "", 0.6, 2.0, 6.1, 4.3, size=12, bullets=[
        "who traded - the market record has no names",
        "data is NSE per-minute summaries from Yahoo Finance - not NSE's own list of every trade; no BSE",
        "personal ‘I bought …’ claims - tested only on 7 made-up examples",
        "that the fixes work on new days - tested only on the days used for fixing",
        "agreement with human experts - right/wrong judged by a separate AI model",
        "what the speaker saw - a TIMING PASS only means the price fits the record",
        "whether SEBI's rule applies - depends on who the speaker is",
        "small samples (dating checked on 3 recordings); one TV chart design; dating searches only the 7 trading days in our data"])
    d.text(s, "What comes next", 6.9, 1.55, 6, 0.4, size=15, color=ACCENT, bold=True)
    d.text(s, "", 6.9, 2.0, 6.0, 2.5, size=12, bullets=[
        "lock the rules, test on new TV days, with human judges",
        "influencer videos with real ‘I bought …’ claims and screenshots of their trades",
        "licensed exchange data listing every trade, and BSE",
        "check whether a price appeared on screen before it was spoken",
        "answer ‘can't tell’ when speech or screen reading is unsure",
        "teach the software other channels' chart designs"])
    d.text(s, "Design trade-offs", 6.9, 4.55, 6, 0.35, size=13, color=WARN, bold=True)
    d.text(s, "", 6.9, 4.9, 6.0, 1.6, size=11, bullets=[
        ("Checking window: ", "wider is safer if the time is uncertain but makes the test weaker; ours = time uncertainty + 3 minutes"),
        ("Heavily traded stocks: ", "far more shares change hands than any claim, so TRADE almost never fails there - a TRADE PASS is weak evidence"),
        ("Targets vs prices now: ", "confusing them is the main cause of wrong ‘impossible’ answers → when unsure, ABSTAIN")])

    # 15 ── conclusion + references
    s = d.slide("Takeaway: a checker for market videos must refuse to guess", "Conclusion and references", section="Outlook", notes=N(
        "Three takeaways: a new checking task; dating a video from its own chart; and the finding that careless checking is wrong almost every "
        "time it accuses. All within the limits on slide 14. Thank you.",
        ("What is proven, what is not?", "Proven: the two guarantees under their conditions; on our data, the date was right whenever given, "
                                         "and careless checking mostly accuses wrongly. Not yet: that the fixes work on new days, and real personal trades.")))
    for i, (h, b) in enumerate([("A new checking task", "for each claim in a video: fits / impossible / can't tell, against the market record - "
                                 "‘impossible’ only when no sensible reading fits"),
                                ("Dating a video from its own chart", "the day's high only rises, so with the price it pins date and time - "
                                 "give or take 4–30 s; the date was right every time one was given"),
                                ("Careful checking is necessary", f"our own careless first version's ‘impossible’ was right only {v1_ok} of {v1_con} times on new TV days")]):
        d.box(s, f"{i + 1}", 0.6, 1.65 + i * 1.3, 0.6, 1.1, size=22, fill=ACCENT, color=WHITE, line=ACCENT)
        d.text(s, h, 1.35, 1.65 + i * 1.3, 5.4, 0.4, size=14, color=ACCENT, bold=True)
        d.text(s, b, 1.35, 2.05 + i * 1.3, 5.35, 0.7, size=12, color=INK2)
    d.box(s, "Not yet measured: the final version on new TV days; real personal trade claims. Also: AI judges, small samples, vendor "
              "minute data (slide 14).", 0.6, 5.6, 6.2, 0.8, size=11, fill=PINK, color=BAD, line=PINK)
    d.text(s, "References (research papers and data sources)", 7.0, 1.55, 5.9, 0.35, size=13, color=ACCENT, bold=True)
    d.text(s, "", 7.0, 1.9, 5.95, 5.1, size=8.5, bullets=[
        "[1] Thorne et al. FEVER: a large-scale dataset for Fact Extraction and VERification. NAACL 2018.",
        "[2] Chen et al. TabFact: a large-scale dataset for table-based fact verification. ICLR 2020.",
        "[3] Aly et al. FEVEROUS: Fact Extraction and VERification Over Unstructured and Structured information. NeurIPS D&B 2021.",
        "[4] Venktesh V. et al. QuanTemp: a real-world open-domain benchmark for fact-checking numerical claims. SIGIR 2024.",
        "[5] Luo, Darrell, Rohrbach. NewsCLIPpings: automatic generation of out-of-context multimodal media. EMNLP 2021.",
        "[6] Geifman, El-Yaniv. Selective classification for deep neural networks. NeurIPS 2017.",
        "[7] Manski. Partial Identification of Probability Distributions. Springer, 2003.",
        "[8] Radford et al. Robust speech recognition via large-scale weak supervision (Whisper). ICML 2023.",
        "[9] Grigoras. Digital audio recording analysis: the Electric Network Frequency (ENF) criterion. Int. J. Speech, Language and the Law 12(1), 2005.",
        "[10] Luo, Li, Wang, Lin. ChartOCR: data extraction from charts images via a deep hybrid framework. WACV 2021.",
        "[11] Liu et al. DePlot: one-shot visual language reasoning by plot-to-table translation. Findings of ACL 2023.",
        "[12] Sakoe, Chiba. Dynamic programming algorithm optimization for spoken word recognition. IEEE Trans. ASSP 26(1), 1978.",
        "[13] SEBI circular HO/47/17/12(11)2025-MRD-POD3/I/11107/2026, 8 May 2026 (in force 1 Jul 2026), FAQ Q8.",
        "[14] Yahoo Finance 1-minute data for NSE: day high/low from 10 Sep 2026, full minute data 29 Sep–8 Oct 2026 (vendor data)."])

    # Appendix A ── all rules
    s = d.slide("Appendix A · The decision rules in symbols (for specialists)", "Appendix", notes=N("Backup: every rule, in words and symbols."))
    table(d, s, [["Rule", "In plain words", "Formal"],
                 ["Time map", "per day, keep shifts where price AND day-high match the record in every frame",
                  "T_d = ⋂_k {τ : p_k ∈ [lo,hi]_d(t_k+τ) ± ε  ∧  RunMax_d(t_k+τ) ∈ [h_k−1, h_k+1]}"],
                 ["Accept time", "one day, one stretch, ≤ ±90 s; else WEAK/ABSTAIN + labelled clock", "IDENTIFIED ⇔ |{d : T_d ≠ ∅}| = 1, T_d one interval, u ≤ 90 s"],
                 ["Compatibility", "stated value (± rounding) inside a minute the speaker could quote", "ok(c,a) ⇔ ∃m ∈ W_a(t+τ): P ∈ [lo_m−δ, hi_m+δ]   (‘above X’: hi_m ≥ X−δ)"],
                 ["FAIL", "complete data, clear current-price claim, every reading fails", "complete(W) ≥ 90 % ∧ clear(c) ∧ ∀a∈𝒜 ¬ok(c,a)   (day high/low: gap > 0.25 %)"],
                 ["PASS / ABSTAIN", "one admissible reading fits / anything else", "PASS ⇔ ∃a ok(c,a);  ABSTAIN otherwise"],
                 ["Trade", "no more shares than all of NSE traded at that price then; holiday = no trading", "FAIL ⇔ holiday ∨ ∀a ¬ok ∨ q > Σ{V_m : P ∈ [lo_m,hi_m]}"],
                 ["Tenure", "age of the data when used; the rule binds only some people and conduct", "PASS ⇔ age ≥ 30 d ∨ use < 1 Jul 2026; else ABSTAIN (never FAIL)"]],
          0.6, 1.6, 12.2, [1.7, 4.8, 5.7], size=11, row_h=0.64)

    # Appendix B ── proofs
    s = d.slide("Appendix B · Proofs of the two guarantees (for specialists)", "Appendix", notes=N("Backup: short proofs and the exact assumptions."))
    d.text(s, "Assumptions", 0.6, 1.6, 12, 0.35, size=14, color=ACCENT, bold=True)
    d.text(s, "", 0.6, 1.95, 12.2, 1.4, size=12, bullets=[
        ("A1 reading: ", "each OCR reading p_k is the displayed price within ε = 0.10; the label h_k equals the day-high so far rounded, "
                         "truncated or ceiled (|h_k − RunMax| ≤ 1)."),
        ("A2 feed: ", "the display and the reference record report the same market (same instrument, same minute bars), with one constant video→market shift τ* within the recording."),
        ("A3 claim: ", "the true reading/assumption a* of claim c lies in 𝒜, and the record is correct on the window W_{a*}.")])
    d.box(s, "Proposition 1 (‘true time kept’). Under A1–A2, the true shift τ* of the true day d* satisfies τ* ∈ T_d*.", 0.6, 3.45, 12.2, 0.5, size=13,
          fill=LIGHT, color=ACCENT, line=LIGHT)
    d.text(s, "Proof. Fix frame k. At market time t_k + τ* on d*, the displayed price lies in that minute's [lo, hi] (A2), so p_k lies in it "
              "within ε (A1); the displayed label is the running maximum up to that time, so RunMax lies in [h_k − 1, h_k + 1] (A1). Hence τ* "
              "belongs to frame k's feasible set S_k. As this holds for every k, τ* ∈ ⋂_k S_k = T_d*.  ∎   (Rejection - no day with T_d ≠ ∅ - "
              "therefore certifies that A1 or A2 failed, and is reported as ABSTAIN.)", 0.6, 3.95, 12.2, 1.1, size=12)
    d.box(s, "Proposition 2 (‘never wrongly impossible’). Under A3, if the verifier outputs FAIL for c, then c is false.", 0.6, 5.15, 12.2, 0.5, size=13, fill=LIGHT, color=ACCENT, line=LIGHT)
    d.text(s, "Proof. FAIL requires ¬ok(c, a) for every a ∈ 𝒜. By A3, a* ∈ 𝒜, so ¬ok(c, a*): under the true reading no minute in the true "
              "window holds P ± δ in the (correct) record, i.e. the stated value did not occur - c is false.  ∎   The converse does not hold: "
              "PASS (∃a ok) does not imply c is true; ABSTAIN covers everything else.", 0.6, 5.65, 12.2, 1.1, size=12)

    # Appendix C ── glossary
    s = d.slide("Appendix C · Glossary - every technical word in this deck", "Appendix", notes=N("Backup: definitions of every technical term."))
    terms = [("PASS / FAIL / ABSTAIN", "fits the record / impossible however the sentence is read / not enough evidence to tell"),
             ("TRADE · TIMING · TENURE", "the three checks; in the app: Fill conservation · Knowability · Data age"),
             ("Market record", "NSE prices per minute (lowest, highest, shares traded), from Yahoo Finance; no names"),
             ("Claim", "any price, quantity or trade stated or shown in the video"),
             ("Reader / encoder · OCR · ASR", "program turning one input into facts · reads text in pictures · turns speech into text"),
             ("Day-high (‘ratchet’, monotone)", "the chart's ‘highest so far today’; it can only rise, so it works like a clock"),
             ("τ (tau), ± u", "how far the video runs behind the market, and the uncertainty of that estimate"),
             ("Feasible set", "all moments still consistent with every frame seen so far"),
             ("Sound / admissible reading", "never says FAIL wrongly / a reading of the sentence a listener could reasonably take"),
             ("Proposition / proof", "a precisely stated guarantee / the argument showing it holds (Appendix B)"),
             ("Held-out · development · frozen", "fair test on unseen data · data used for fixing (weaker) · locked, no changes"),
             ("Baseline · ablation", "simpler method compared against · switching one part off to show it is needed"),
             ("95 % range (Wilson interval)", "with few cases, the range in which the true rate very likely lies; e.g. 3 of 3 right → somewhere between 44 % and 100 %"),
             ("ENF · DTW", "dating audio from electric-grid hum · lining up two time series"),
             ("SNTP clock · IST", "the recording computer's clock, set from internet time servers - used only to check · Indian Standard Time"),
             ("Fingerprint (hash, sha256)", "a short code computed from a file; it changes if the file is altered"),
             ("Liquid stock · option strike", "a heavily traded stock · the fixed price in an option contract (‘the 22,700 call’)"),
             ("SEBI 30-day rule", "a person who only teaches about stocks must not use prices from the last 30 days when recommending a stock")]
    half = (len(terms) + 1) // 2
    for col, chunk in enumerate((terms[:half], terms[half:])):
        table(d, s, [["Term", "Meaning"]] + [[a, b] for a, b in chunk], 0.6 + col * 6.2, 1.6, 6.0, [2.1, 3.9], size=10, row_h=0.5)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    d.p.save(OUT)
    lines = ["# Presenter guide - Did That Trade Really Happen? (Capstone Review 2)", "",
             "Generated by `python -m scripts.build_ppt` from the speaker notes. Read the script aloud; use the Q&A if asked.", "",
             "## Words you will hear (glossary)", ""]
    lines += [f"- **{w}**: {m}" for w, m in GLOSSARY]
    for i, sl in enumerate(d.p.slides, 1):
        texts = [sh.text_frame.text for sh in sl.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
        title = next((t for t in texts if len(t) > 12 and not t.isupper()), texts[0])
        lines += ["", f"## Slide {i}: {title}", "", sl.notes_slide.notes_text_frame.text.replace("SAY:\n", "").replace(
            "IF THE PANEL ASKS:", "**If the panel asks:**").replace("\nQ: ", "\n- **Q:** ").replace("\nA: ", "\n  **A:** ")]
    (OUT.parent / "PRESENTER_GUIDE.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(OUT, d.n, "slides")


if __name__ == "__main__":
    main()
